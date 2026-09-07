"""The deck must speak the language its source document speaks.

The template is a Japanese deck. Rendering an English proposal through it used
to produce a bilingual result — English body copy under Japanese headings, and
a cover reading `○○株式会社　御中` above the client's real name. These render
for real and read the finished file back, because that mixture was only ever
visible there.

The Japanese path is asserted alongside every English one: none of this may
change what a Japanese proposal renders to.
"""
from __future__ import annotations

import re

import pytest

from lib.profile_parser import parse_profile
from lib.renderer import PPTXRenderer

pptx = pytest.importorskip('pptx')

# Kana only. The Sun* copyright footer carries 株式会社 on every slide of every
# deck — it is a legal entity name, not chrome, and is deliberately left alone.
KANA = re.compile(r'[぀-ゟ゠-ヿ]')


def _walk(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def _render(fixtures_dir, template_path, tmp_path, name):
    profile = parse_profile((fixtures_dir / f'{name}.md').read_text(encoding='utf-8'))
    out = PPTXRenderer().render_from_profile(
        profile, str(template_path), name, str(tmp_path),
        extra_slides=list(profile.auto_extra_slides) or None,
    )
    return pptx.Presentation(out)


def _texts(prs):
    return [shape.text_frame.text.strip()
            for slide in prs.slides
            for shape in _walk(slide.shapes)
            if shape.has_text_frame and shape.text_frame.text.strip()]


@pytest.fixture(scope='module')
def english_deck(fixtures_dir, template_path, tmp_path_factory):
    return _render(fixtures_dir, template_path,
                   tmp_path_factory.mktemp('en'), 'english-proposal')


@pytest.fixture(scope='module')
def japanese_deck(fixtures_dir, template_path, tmp_path_factory):
    return _render(fixtures_dir, template_path, tmp_path_factory.mktemp('ja'),
                   'generic-jp-with-sections-mapping')


# --- source language ---------------------------------------------------------

@pytest.mark.parametrize('name, expected', [
    ('english-proposal', False),
    ('generic-jp-with-sections-mapping', True),
])
def test_source_language_is_detected(fixtures_dir, name, expected):
    profile = parse_profile((fixtures_dir / f'{name}.md').read_text(encoding='utf-8'))
    assert profile.is_japanese_source is expected


# --- cover -------------------------------------------------------------------

def test_english_cover_drops_the_jp_honorific_line(english_deck):
    """`○○株式会社　御中` is a placeholder AND an honorific — neither survives."""
    cover = _texts(english_deck)[:3]
    assert '御中' not in ' '.join(cover)
    assert 'Evering' in cover
    assert '2026.04.19' in cover


def test_japanese_cover_keeps_the_honorific_line(japanese_deck):
    cover = ' '.join(_texts(japanese_deck)[:3])
    assert '御中' in cover
    assert '東京物流株式会社' in cover


# --- slide headings ----------------------------------------------------------

def test_slide_is_titled_from_the_source_heading(english_deck):
    """Breadcrumb tail and the title below it both come from the md heading.

    The template says `System Overview ｜ プロジェクト背景` / `プロジェクト背景`
    — the chapter name of the proposal it was authored for.
    """
    texts = _texts(english_deck)
    assert any(t.endswith('｜  Project Background') for t in texts), texts[:20]
    assert 'Project Background' in texts
    assert 'プロジェクト背景' not in texts


def test_title_spelled_differently_from_its_breadcrumb_is_still_retitled(english_deck):
    """Slide 5's breadcrumb says `機能一覧`, its title says `機能一覧表`.

    An exact-match rule caught the breadcrumb and left the title in Japanese.
    """
    texts = _texts(english_deck)
    assert 'Feature List' in texts
    assert '機能一覧表' not in texts


def test_japanese_headings_come_from_the_japanese_source(japanese_deck):
    """Same pass, Japanese document: the SOURCE's wording, not the template's."""
    texts = _texts(japanese_deck)
    assert '現行業務の課題と狙い' in texts        # the fixture's own heading
    assert 'プロジェクト背景' not in texts        # the template's


# --- fixed chrome ------------------------------------------------------------

def test_english_deck_ships_no_japanese_chrome(english_deck):
    """Column headers and captions no heading can name are Englished too."""
    leftovers = [t for t in _texts(english_deck) if KANA.search(t)]
    assert leftovers == [], leftovers


def test_japanese_deck_keeps_its_chrome_verbatim(japanese_deck):
    texts = _texts(japanese_deck)
    assert '現状の課題' in texts
    assert 'Current Issues' not in texts
