"""Pins `_LATIN_EM_WIDTH` and `_COVER_FIT_MARGIN` to observed behaviour.

`renderer.py` documents these two constants as "calibrated against rendered
output rather than nominal font metrics" and pointed at a `_COVER_FIT_CASES`
table — which no longer existed anywhere, leaving the numbers unpinned. This is
that table, rebuilt against the shipped template's real cover geometry.

The width model only drives a WARNING, never a change to the deck, so the bar is
"separates a client name from a proposal title", not exact typography.
"""
from __future__ import annotations

import pytest

from lib.renderer import PPTXRenderer


# Measured from `templates/SVN Proposal Menu.pptx`, slide 1, the shape
# `COVER_SLIDE_LAYOUT['company']` addresses: 6.979in wide, 33pt bold.
COVER_BOX_WIDTH_IN = 6.979
COVER_FONT_PT = 33.0


@pytest.mark.parametrize('text, should_fit', [
    # Real client names — every one of these must ship without a warning.
    ('Styleport', True),
    ('東京物流株式会社', True),
    ('株式会社スタイルポート', True),
    # `cover-and-agenda.md` names this exact string as what must NEVER go in
    # the box: it is a proposal title, not a company, and it wraps over the
    # date line below. The model has to catch it.
    ('ROOV compass 住戸検索 開発', False),
    # A genuinely over-long company name still warns rather than silently
    # overlapping.
    ('株式会社スタイルポートホールディングス', False),
])
def test_cover_fit_cases(text, should_fit):
    assert PPTXRenderer._cover_fits(text, COVER_FONT_PT, COVER_BOX_WIDTH_IN) is should_fit


def test_cover_box_geometry_matches_the_template(template_path):
    """The cases above are only meaningful while the box is this size."""
    from pptx import Presentation

    from lib.templates.svn import COVER_SLIDE_LAYOUT

    shape = Presentation(str(template_path)).slides[0].shapes[COVER_SLIDE_LAYOUT['company']]
    width_in = shape.width / PPTXRenderer._EMU_PER_INCH
    font_pt = next((run.font.size.pt
                    for para in shape.text_frame.paragraphs
                    for run in para.runs if run.font.size), None)
    assert width_in == pytest.approx(COVER_BOX_WIDTH_IN, abs=0.01)
    assert font_pt == pytest.approx(COVER_FONT_PT, abs=0.01)


def test_cjk_counts_wider_than_latin():
    """The model's one structural claim: CJK sets ~1em, latin a fraction."""
    cjk = PPTXRenderer._text_width_in('あああああ', 10)
    latin = PPTXRenderer._text_width_in('aaaaa', 10)
    assert cjk > latin
    assert latin == pytest.approx(cjk * PPTXRenderer._LATIN_EM_WIDTH, rel=1e-6)
