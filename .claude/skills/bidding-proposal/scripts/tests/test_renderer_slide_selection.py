"""Regressions for renderer defects that put the WRONG SLIDE in a deck.

These render for real and read the result back, because the bugs were only
visible in the finished file — the parser output was correct in every case.
"""
from __future__ import annotations

import pytest

from lib.profile_parser import parse_profile
from lib.renderer import PPTXRenderer

pptx = pytest.importorskip('pptx')


def render(md_path, template_path, tmp_path, name='out'):
    profile = parse_profile(md_path.read_text(encoding='utf-8'))
    renderer = PPTXRenderer()
    out = renderer.render_from_profile(
        profile, str(template_path), name, str(tmp_path),
        extra_slides=list(profile.auto_extra_slides) or None,
    )
    return pptx.Presentation(out)


def _walk(shapes):
    """Every shape, descending into groups — the template nests titles in them."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)
        else:
            yield shape


def slide_texts(prs):
    """Flat text of every slide, one string per slide.

    Table cells are included deliberately: most of this template's content sits
    in tables, not text frames, so a text-frame-only reader reports a filled
    slide as blank — which is exactly the false negative that makes these
    assertions look like code failures.
    """
    out = []
    for slide in prs.slides:
        parts = []
        for shape in _walk(slide.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            elif getattr(shape, 'has_table', False) and shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        out.append(' '.join(parts))
    return out


# --- B1: divider addressed by template number after overflow shifted the deck --

def test_divider_matches_chapter_under_overflow(fixtures_dir, template_path, tmp_path):
    """A 40-row table inserts 5 continuation slides before the divider lookup.

    `_dividers_by_label` read `prs.slides[slide_no - 1]` — template numbering —
    and so pulled slide 22 (お見積りの前提条件) while believing it was 17
    (Sun*のスコープ). The deck shipped an Estimated Assumptions divider for a
    document with no such chapter.
    """
    prs = render(fixtures_dir / 'overflow-forces-divider-shift.md',
                 template_path, tmp_path)
    texts = slide_texts(prs)
    assert any('Sun*のスコープ' in t for t in texts)
    assert not any('お見積りの前提条件' in t for t in texts), \
        'the 前提条件 divider belongs to a chapter this document does not have'


# --- S5: agenda label matched by bare prefix --------------------------------

@pytest.mark.parametrize('title, expected_label', [
    ('お見積りの前提条件', ''),                       # ≠ お見積り: different chapter
    ('費用対効果の分析', ''),                         # ≠ 費用
    ('Sun*のスコープ', 'Sun* のスコープ・役割分担'),    # extension at `・` — same chapter
    ('お見積り', 'お見積り'),                          # exact
])
def test_agenda_match_requires_exact_or_separator_boundary(title, expected_label):
    pairs = [('お見積り', 'Pricing'),
             ('Sun* のスコープ・役割分担', 'Scope & Responsibility Split'),
             ('費用', 'Cost')]
    label, _ = PPTXRenderer._agenda_match(title, pairs)
    assert label == expected_label


# --- Phase 07: unfilled slides in a multi-slide section ship template samples --

def test_unfilled_section_slides_are_dropped(template_path, tmp_path):
    """assumptions owns slides 23/24/25 via `[0:5]`, `[5:8]`, `[8:9]`.

    One assumption fills only slide 23; 24 and 25 stayed in the deck holding the
    TEMPLATE'S sample preconditions (デザイン / インフラ), which read as this
    proposal's own contractual terms.
    """
    md = tmp_path / 'sparse.md'
    md.write_text(
        '# T\n\n## Assumptions\n\n### 開発方針\n\nSingle tenant only.\n',
        encoding='utf-8',
    )
    texts = ' '.join(slide_texts(render(md, template_path, tmp_path)))
    assert 'Single tenant only.' in texts
    assert 'A案の場合' not in texts, 'template sample assumption shipped as real content'
    assert 'UIに関しては既存API' not in texts


def test_filled_section_slides_are_all_kept(template_path, tmp_path):
    """The drop must be driven by observed fill, not by a guess about counts."""
    rows = '\n'.join(f'### L{i}\n\nC{i}\n' for i in range(1, 10))
    md = tmp_path / 'full.md'
    md.write_text(f'# T\n\n## Assumptions\n\n{rows}', encoding='utf-8')
    texts = ' '.join(slide_texts(render(md, template_path, tmp_path)))
    for i in (1, 6, 9):  # one from each of the [0:5] / [5:8] / [8:9] slices
        assert f'C{i}' in texts, f'assumption {i} lost — a filled slide was dropped'
