#!/usr/bin/env python3
"""verify-template-contract: assert svn.py's template mappings still resolve.

`lib/templates/svn.py` addresses the SVN Proposal template by exact string
(`SLIDE_TEXT_OVERRIDES`, `SLIDE_TEXT_FROM_PROFILE`, `TEMPLATE_NOTE_TEXTS`) and
by shape index (`COVER_SLIDE_LAYOUT`, `AGENDA_SLIDE_LAYOUT`,
`CHAPTER_DIVIDER_SLIDES`, `EXTRA_SLIDE_CHROME_SLIDE`). That precision is
deliberate — a fuzzy rule would delete real content — but it means editing the
.pptx silently disables a mapping instead of erroring.

This script is the gate. Run it after ANY change to the template file; a
non-zero exit means a mapping no longer describes the deck it addresses.

    $ verify-template-contract.py
    [OK]   SLIDE_TEXT_OVERRIDES     4/4 matched
    [FAIL] TEMPLATE_NOTE_TEXTS     12/13 matched
           unmatched: 'Kế hoạch thực hiện'

A failure is reported, never "fixed" by deleting the entry: dropping a
TEMPLATE_NOTE_TEXTS line re-ships a Vietnamese authoring note to the client.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from ensure_deps import ensure_deps
ensure_deps()

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from lib.templates.svn import (
    AGENDA_SLIDE_LAYOUT,
    CHAPTER_DIVIDER_SLIDES,
    COVER_JP_HONORIFIC_IDX,
    COVER_SLIDE_LAYOUT,
    EXTRA_SLIDE_CHROME_SLIDE,
    SLIDE_TEXT_FROM_PROFILE,
    SLIDE_TEXT_OVERRIDES,
    TEMPLATE_CHROME_EN,
    TEMPLATE_NOTE_TEXTS,
)

DEFAULT_TEMPLATE = 'SVN Proposal Menu.pptx'


def walk_shapes(shapes):
    """Yield every shape, descending into groups — mirrors the renderer."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)
        else:
            yield shape


def normalize_note_text(text: str) -> str:
    """Same normalization `PPTXRenderer._normalize_note_text` applies."""
    lines = [line.strip() for line in (text or '').splitlines()]
    return '\n'.join(line for line in lines if line)


def slide_texts(slide, *, normalized: bool = False) -> set:
    out = set()
    for shape in walk_shapes(slide.shapes):
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        out.add(normalize_note_text(text) if normalized else text.strip())
    return out


class Report:
    """Collects per-check results and renders the summary."""

    def __init__(self):
        self.failed = False

    def check(self, name: str, total: int, unmatched: list, detail: str = '') -> None:
        ok = not unmatched
        self.failed |= not ok
        tag = '[OK]  ' if ok else '[FAIL]'
        print(f'{tag} {name:26} {total - len(unmatched)}/{total} {detail or "matched"}')
        for item in sorted(str(i) for i in unmatched):
            print(f'       unmatched: {item}')


def check_slide_text_map(report: Report, prs, name: str, mapping: dict,
                         after: dict | None = None) -> None:
    """Every key text must exist on the slide its entry names.

    `after` is a map already applied to the deck before this one runs at render
    time. The two chain: slide 39's title is `想定成果物（1/2）- Advanced` in the
    file, SLIDE_TEXT_OVERRIDES rewrites it to `想定成果物`, and only then does
    SLIDE_TEXT_FROM_PROFILE key off that. Checking the raw template would call
    the second map stale when it is simply downstream of the first.
    """
    total = sum(len(texts) for texts in mapping.values())
    unmatched = []
    for slide_no, texts in mapping.items():
        if slide_no > len(prs.slides._sldIdLst):
            unmatched.extend(f'slide {slide_no} (out of range): {t!r}' for t in texts)
            continue
        present = slide_texts(prs.slides[slide_no - 1])
        for original, replacement in (after or {}).get(slide_no, {}).items():
            if original in present:
                present.discard(original)
                present.add(replacement)
        unmatched.extend(f'slide {slide_no}: {t!r}' for t in texts if t not in present)
    report.check(name, total, unmatched)


def check_note_texts(report: Report, prs) -> None:
    """Every note text must appear somewhere in the deck."""
    present = set()
    for slide in prs.slides:
        present |= slide_texts(slide, normalized=True)
    unmatched = [t for t in TEMPLATE_NOTE_TEXTS if t not in present]
    report.check('TEMPLATE_NOTE_TEXTS', len(TEMPLATE_NOTE_TEXTS), unmatched)


def check_divider_slides(report: Report, prs) -> None:
    """Each divider must exist and hold text at the EN and JP shape indices."""
    unmatched = []
    for slide_no in CHAPTER_DIVIDER_SLIDES:
        if slide_no > len(prs.slides._sldIdLst):
            unmatched.append(f'slide {slide_no}: out of range')
            continue
        shapes = prs.slides[slide_no - 1].shapes
        if len(shapes) <= 1:
            unmatched.append(f'slide {slide_no}: fewer than 2 shapes')
            continue
        for idx, role in ((0, 'EN'), (1, 'JP')):
            if not (shapes[idx].has_text_frame and shapes[idx].text_frame.text.strip()):
                unmatched.append(f'slide {slide_no}: shape {idx} ({role}) holds no text')
    report.check('CHAPTER_DIVIDER_SLIDES', len(CHAPTER_DIVIDER_SLIDES) * 2,
                 unmatched, 'shape slots resolve')


def check_shape_layout(report: Report, prs, name: str, slide_no: int, layout: dict) -> None:
    """Each mapped index must be in range on its slide and hold a text frame."""
    unmatched = []
    if slide_no > len(prs.slides._sldIdLst):
        unmatched = [f'slide {slide_no}: out of range']
    else:
        shapes = prs.slides[slide_no - 1].shapes
        for field, idx in layout.items():
            if idx >= len(shapes):
                unmatched.append(f'{field}: index {idx} past {len(shapes)} shapes')
            elif not shapes[idx].has_text_frame:
                unmatched.append(f'{field}: shape {idx} has no text frame')
    report.check(name, len(layout), unmatched, f'indices resolve on slide {slide_no}')


def check_chrome_labels(report: Report, prs) -> None:
    """Every Englished label must still exist somewhere in the deck.

    An entry matching nothing is dead weight, and dead weight is how a label
    silently starts shipping in Japanese again after the template rewords it.
    """
    present = set()
    for slide in prs.slides:
        present |= slide_texts(slide)
    unmatched = [t for t in TEMPLATE_CHROME_EN if t not in present]
    report.check('TEMPLATE_CHROME_EN', len(TEMPLATE_CHROME_EN), unmatched)


def check_cover_honorific(report: Report, prs) -> None:
    """The shape a non-Japanese deck removes must be the honorific line.

    Addressed by index, and the index is only meaningful while that shape is
    what it was. If the template reorders slide 1, this check fails here rather
    than deleting the client's name from a deck.
    """
    unmatched = []
    shapes = prs.slides[0].shapes
    if COVER_JP_HONORIFIC_IDX >= len(shapes):
        unmatched.append(f'index {COVER_JP_HONORIFIC_IDX} past {len(shapes)} shapes')
    else:
        shape = shapes[COVER_JP_HONORIFIC_IDX]
        text = shape.text_frame.text.strip() if shape.has_text_frame else ''
        if '御中' not in text:
            unmatched.append(f'shape {COVER_JP_HONORIFIC_IDX} is {text!r}, not the '
                             f'御中 line')
    report.check('COVER_JP_HONORIFIC_IDX', 1, unmatched,
                 'the honorific line is where svn.py says')


def check_chrome_slide(report: Report, prs) -> None:
    unmatched = []
    idx = EXTRA_SLIDE_CHROME_SLIDE - 1
    if not (0 <= idx < len(prs.slides)):
        unmatched.append(f'slide {EXTRA_SLIDE_CHROME_SLIDE}: out of range')
    elif prs.slides[idx].slide_layout is None:
        unmatched.append(f'slide {EXTRA_SLIDE_CHROME_SLIDE}: no slide layout')
    report.check('EXTRA_SLIDE_CHROME_SLIDE', 1, unmatched,
                 f'slide {EXTRA_SLIDE_CHROME_SLIDE} usable as chrome source')


def resolve_template(name: str) -> Path:
    """Resolve the template the same way gen-slide.py does.

    Kept inside the skill directory on purpose — this is a contract check
    against the template that SHIPS with the skill, not a general .pptx linter.
    """
    candidate = Path(name)
    if candidate.is_absolute():
        return candidate
    here = Path(__file__).parent
    for base in (here / 'templates' / name, here.parent / 'templates' / name):
        if base.exists():
            return base
    return candidate


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--template', default=DEFAULT_TEMPLATE,
                    help=f'Template PPTX name or path (default: {DEFAULT_TEMPLATE})')
    args = ap.parse_args()

    template = resolve_template(args.template)
    if not template.exists():
        print(f'[FAIL] template not found: {template}')
        return 1
    print(f'Template: {template}')

    prs = Presentation(str(template))
    report = Report()
    check_slide_text_map(report, prs, 'SLIDE_TEXT_OVERRIDES', SLIDE_TEXT_OVERRIDES)
    check_slide_text_map(report, prs, 'SLIDE_TEXT_FROM_PROFILE', SLIDE_TEXT_FROM_PROFILE,
                         after=SLIDE_TEXT_OVERRIDES)
    check_note_texts(report, prs)
    check_divider_slides(report, prs)
    check_shape_layout(report, prs, 'COVER_SLIDE_LAYOUT', 1, COVER_SLIDE_LAYOUT)
    check_shape_layout(report, prs, 'AGENDA_SLIDE_LAYOUT', 2, AGENDA_SLIDE_LAYOUT)
    check_cover_honorific(report, prs)
    check_chrome_labels(report, prs)
    check_chrome_slide(report, prs)

    if report.failed:
        print('\nA mapping no longer describes the shipped template. Report it — '
              'do NOT delete the entry to make this pass: dropping a note entry '
              're-ships template authoring guidance to the client.')
        return 1
    print('\nAll template mappings resolve.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
